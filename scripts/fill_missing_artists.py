from __future__ import annotations

import argparse
import csv
import os
import shutil
import sys
import tempfile
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from tag_creator.missing_artist import (  # noqa: E402
    MissingArtistResolver,
    SUPPORTED_TABULAR_EXTENSIONS,
    build_header_map,
    clean_text,
    filename_from_row,
    is_missing_artist_value,
    is_missing_title_value,
    read_tabular_rows,
    repair_rows,
    row_value,
    write_tabular_rows,
)


def output_copy_path(path: Path, suffix: str) -> Path:
    return path.with_name(f"{path.stem}{suffix}{path.suffix}")


def find_tabular_files(paths: list[Path], recursive: bool) -> list[Path]:
    files: list[Path] = []
    for path in paths:
        if path.is_file() and path.suffix.lower() in SUPPORTED_TABULAR_EXTENSIONS:
            files.append(path)
            continue
        if not path.is_dir():
            continue
        pattern = "**/*" if recursive else "*"
        files.extend(
            candidate
            for candidate in path.glob(pattern)
            if candidate.is_file()
            and candidate.suffix.lower() in SUPPORTED_TABULAR_EXTENSIONS
            and not candidate.name.startswith("~$")
            and "_artist_fixed" not in candidate.stem
        )
    return sorted(set(files))


def repair_xlsx_file(
    path: Path,
    target_path: Path,
    *,
    resolver: MissingArtistResolver,
    remove_unresolved: bool,
) -> tuple[int, int, int, int]:
    from openpyxl import load_workbook

    if target_path != path:
        target_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(path, target_path)

    workbook = load_workbook(target_path)
    sheet = workbook[workbook.sheetnames[0]]
    headers = [clean_text(cell.value) for cell in sheet[1]]
    header_map = build_header_map(headers)
    if "artist" not in header_map:
        artist_col = len(headers) + 1
        sheet.cell(row=1, column=artist_col).value = "artist"
        headers.append("artist")
        header_map = build_header_map(headers)
    artist_col = headers.index(header_map["artist"]) + 1

    checked = filled = removed_identity = removed_unresolved = 0
    rows_to_delete: list[int] = []
    csv_context = target_path.stem.replace("_with_tag", "")

    for excel_row in range(2, sheet.max_row + 1):
        row = {
            headers[index]: clean_text(sheet.cell(row=excel_row, column=index + 1).value)
            for index in range(len(headers))
            if headers[index]
        }
        if not any(row.values()):
            continue
        checked += 1
        title = row_value(row, header_map, "title")
        artist = row_value(row, header_map, "artist")
        if is_missing_title_value(title) and is_missing_artist_value(artist):
            rows_to_delete.append(excel_row)
            removed_identity += 1
            continue
        if not is_missing_artist_value(artist):
            continue

        resolution = resolver.resolve(
            title=title,
            filename=filename_from_row(row, header_map),
            file_path=row_value(row, header_map, "file_path", "path", "filepath"),
            existing_artist=artist,
            row=row,
            header_map=header_map,
            csv_context=csv_context,
        )
        if resolution and resolution.artist:
            sheet.cell(row=excel_row, column=artist_col).value = resolution.artist
            filled += 1
        elif remove_unresolved:
            rows_to_delete.append(excel_row)
            removed_unresolved += 1

    for excel_row in sorted(rows_to_delete, reverse=True):
        sheet.delete_rows(excel_row, 1)

    handle = tempfile.NamedTemporaryFile(delete=False, dir=str(target_path.parent), suffix=".xlsx")
    temp_path = Path(handle.name)
    handle.close()
    try:
        workbook.save(temp_path)
        temp_path.replace(target_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()
    return checked, filled, removed_identity, removed_unresolved


def repair_csv_like_file(
    path: Path,
    target_path: Path,
    *,
    resolver: MissingArtistResolver,
    remove_unresolved: bool,
) -> tuple[int, int, int, int]:
    headers, rows = read_tabular_rows(path)
    repaired_rows, repaired_headers, stats = repair_rows(
        rows,
        headers,
        resolver=resolver,
        csv_context=target_path.stem.replace("_with_tag", ""),
        remove_unresolved=remove_unresolved,
    )
    write_tabular_rows(target_path, repaired_headers, repaired_rows)
    return stats.checked, stats.filled, stats.removed_missing_identity, stats.removed_unresolved_artist


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Fill only missing artist cells in generated tag_creator CSV/XLSX files."
    )
    parser.add_argument(
        "--input",
        action="append",
        default=[],
        help="File or directory to repair. Can be passed multiple times. Defaults to output and clean.",
    )
    parser.add_argument(
        "--media-root",
        action="append",
        default=[],
        help="Mounted media root used to read embedded artist tags. Can be passed multiple times.",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Update files in place. Without this, writes *_artist_fixed copies.",
    )
    parser.add_argument(
        "--keep-unresolved",
        action="store_true",
        help="Keep rows where artist cannot be found after local evidence and Google search.",
    )
    parser.add_argument(
        "--recursive",
        action="store_true",
        default=True,
        help="Scan input directories recursively. Default: true.",
    )
    parser.add_argument(
        "--suffix",
        default="_artist_fixed",
        help="Suffix used for copy mode. Default: _artist_fixed.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    inputs = [Path(value) for value in args.input] or [Path(os.getenv("OUTPUT_DIR", "output")), Path("clean")]
    files = find_tabular_files(inputs, args.recursive)
    if not files:
        print(f"No CSV/XLSX files found in: {', '.join(str(path) for path in inputs)}")
        return 0

    resolver = MissingArtistResolver(media_roots=[Path(root) for root in args.media_root])
    total_checked = total_filled = total_removed_identity = total_removed_unresolved = 0
    processed = 0
    for path in files:
        target_path = path if args.overwrite else output_copy_path(path, args.suffix)
        try:
            if path.suffix.lower() == ".xlsx":
                checked, filled, removed_identity, removed_unresolved = repair_xlsx_file(
                    path,
                    target_path,
                    resolver=resolver,
                    remove_unresolved=not args.keep_unresolved,
                )
            else:
                checked, filled, removed_identity, removed_unresolved = repair_csv_like_file(
                    path,
                    target_path,
                    resolver=resolver,
                    remove_unresolved=not args.keep_unresolved,
                )
        except Exception as exc:
            print(f"failed: {path}\n{exc}")
            return 2

        processed += 1
        total_checked += checked
        total_filled += filled
        total_removed_identity += removed_identity
        total_removed_unresolved += removed_unresolved
        print(
            f"repaired: {target_path} "
            f"(checked={checked}, filled={filled}, "
            f"removed_missing_title_artist={removed_identity}, "
            f"removed_unresolved_artist={removed_unresolved})"
        )

    print(
        "done. "
        f"files={processed}, checked={total_checked}, filled={total_filled}, "
        f"removed_missing_title_artist={total_removed_identity}, "
        f"removed_unresolved_artist={total_removed_unresolved}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

