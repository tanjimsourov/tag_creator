from __future__ import annotations

import csv
import json
import os
import re
import tempfile
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Mapping, Sequence

import requests
from bs4 import BeautifulSoup
from mutagen import File as MutagenFile

from .matching import normalize_text, similarity


MEDIA_EXTENSIONS = {".mp3", ".mp4", ".m4a", ".aac", ".flac", ".wav", ".wma", ".ogg"}
SUPPORTED_TABULAR_EXTENSIONS = {".csv", ".xls", ".xlsx"}
MISSING_ARTIST_VALUES = {
    "",
    "-",
    "n/a",
    "na",
    "none",
    "null",
    "unknown",
    "unknown artist",
    "artist unknown",
    "needs_review",
    "needs review",
    "not listed",
    "not listed in free sources",
}
ARTIST_NOISE_WORDS = {
    "album",
    "audio",
    "cover",
    "download",
    "full song",
    "genre",
    "google",
    "hd",
    "instrumental",
    "karaoke",
    "lyrics",
    "music",
    "official",
    "official audio",
    "official music video",
    "official video",
    "remaster",
    "song",
    "soundtrack",
    "track",
    "video",
    "youtube",
}
TITLE_NOISE_PATTERN = re.compile(
    r"\b(?:official|music|video|audio|lyrics?|karaoke|instrumental|visualizer|hd|4k|"
    r"no guide melody|cover|remaster(?:ed)?|from .+? soundtrack)\b",
    re.IGNORECASE,
)
TITLE_SEPARATOR_PATTERN = re.compile(r"\s+(?:-|--|/|\||:)\s+")
SEARCH_RESULT_SELECTORS = (
    "div.g",
    "div[data-header-feature]",
    "div.MjjYud",
    "div.SoaBEf",
    "li.b_algo",
    "div.tF2Cxc",
)
TRUSTED_SOURCE_HINTS = (
    "apple music",
    "spotify",
    "deezer",
    "musicbrainz",
    "shazam",
    "genius",
    "last.fm",
    "officialcharts",
    "youtube",
)
PROVIDER_ARTIST_COLUMNS = (
    "verified_artist",
    "provider_artist",
    "ai_artist",
    "musicbrainz_artist",
    "spotify_artist",
    "itunes_artist",
    "deezer_artist",
    "lastfm_artist",
    "discogs_artist",
    "acoustid_artist",
    "web_artist",
    "artist_name",
    "artists",
    "track_artist",
    "primary_artist",
)
PROVIDER_JSON_COLUMNS = (
    "merged_json",
    "providers_json",
    "provider_results_json",
    "provider_results",
    "metadata_json",
    "analysis_json",
)
ARTIST_JSON_FIELDS = {"artist", "artists", "artist_name", "track_artist", "primary_artist"}


@dataclass(frozen=True)
class ArtistResolution:
    artist: str
    confidence: float
    source: str
    detail: str = ""


@dataclass
class ArtistRepairStats:
    checked: int = 0
    filled: int = 0
    removed_missing_identity: int = 0
    removed_unresolved_artist: int = 0
    unchanged: int = 0
    local_hits: int = 0
    google_hits: int = 0
    failed_files: int = 0
    files_processed: int = 0
    messages: list[str] = field(default_factory=list)

    def extend(self, other: "ArtistRepairStats") -> None:
        self.checked += other.checked
        self.filled += other.filled
        self.removed_missing_identity += other.removed_missing_identity
        self.removed_unresolved_artist += other.removed_unresolved_artist
        self.unchanged += other.unchanged
        self.local_hits += other.local_hits
        self.google_hits += other.google_hits
        self.failed_files += other.failed_files
        self.files_processed += other.files_processed
        self.messages.extend(other.messages)


def clean_text(value: object) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\u266a", " ").replace("\u266b", " ")
    text = re.sub(r"[\u200b-\u200f\ufeff]", "", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip().strip("\"'")


def normalize_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", clean_text(value).lower()).strip("_")


def build_header_map(headers: Sequence[str]) -> dict[str, str]:
    return {normalize_header(header): header for header in headers if clean_text(header)}


def row_value(row: Mapping[str, object], header_map: Mapping[str, str], *names: str) -> str:
    for name in names:
        actual = header_map.get(normalize_header(name), name)
        if actual in row:
            value = clean_text(row.get(actual, ""))
            if value:
                return value
    return ""


def is_missing_artist_value(value: object) -> bool:
    cleaned = clean_text(value)
    normalized = normalize_text(cleaned)
    return normalized in MISSING_ARTIST_VALUES or normalized in ARTIST_NOISE_WORDS


def is_missing_title_value(value: object) -> bool:
    cleaned = clean_text(value)
    normalized = normalize_text(cleaned)
    return normalized in MISSING_ARTIST_VALUES or normalized in {"title", "song title"}


def _strip_extension_noise(value: str) -> str:
    stem = Path(value).stem if Path(value).suffix.lower() in MEDIA_EXTENSIONS else value
    stem = re.sub(r"[_]+", " ", stem)
    stem = re.sub(r"\s+", " ", stem)
    return stem.strip()


def clean_title_for_search(value: str) -> str:
    title = _strip_extension_noise(clean_text(value))
    title = re.sub(r"\[[^\]]*\]|\([^\)]*\)", " ", title)
    title = TITLE_NOISE_PATTERN.sub(" ", title)
    title = re.sub(r"\s+", " ", title)
    return title.strip(" -_|:")


def _artist_candidate(value: object, *, title: str = "", parsed_title: str = "") -> str:
    candidate = clean_text(value)
    if not candidate:
        return ""
    candidate = re.sub(r"\[[^\]]*\]|\([^\)]*\)", " ", candidate)
    for known_title in (title, parsed_title):
        if known_title:
            title_match = re.search(rf"\b{re.escape(known_title)}\b", candidate, flags=re.IGNORECASE)
            if title_match and title_match.start() > 0:
                candidate = candidate[: title_match.start()]
    candidate = TITLE_NOISE_PATTERN.sub(" ", candidate)
    candidate = re.split(r"\s+-\s+Topic\b", candidate, maxsplit=1, flags=re.IGNORECASE)[0]
    candidate = re.split(r"\s+\b(?:from|on|album|single|released|recorded)\b\s+", candidate, maxsplit=1, flags=re.IGNORECASE)[0]
    candidate = re.sub(r"\b(?:feat|ft|featuring)\.?\b.*$", "", candidate, flags=re.IGNORECASE)
    candidate = re.sub(r"\s+", " ", candidate).strip(" -_|:,.\"'")
    normalized = normalize_text(candidate)
    if (
        not candidate
        or len(candidate) > 80
        or normalized in MISSING_ARTIST_VALUES
        or normalized in ARTIST_NOISE_WORDS
        or normalized == normalize_text(title)
        or normalized == normalize_text(parsed_title)
        or "http" in normalized
        or "/" in candidate
    ):
        return ""
    if len(candidate) <= 1:
        return ""
    return candidate


def parse_artist_title(value: str) -> tuple[str, str]:
    cleaned = _strip_extension_noise(value)
    if not cleaned:
        return "", ""
    for separator in (" - ", " -- ", " | ", " / "):
        if separator in cleaned:
            left, right = cleaned.split(separator, 1)
            artist = _artist_candidate(left)
            title = clean_title_for_search(right)
            if artist and title:
                return artist, title
    return "", ""


def filename_from_row(row: Mapping[str, object], header_map: Mapping[str, str]) -> str:
    filename = row_value(row, header_map, "filename", "file_name", "name")
    if filename:
        return filename
    file_path = row_value(row, header_map, "file_path", "path", "filepath")
    return Path(file_path).name if file_path else ""


def _title_variants(title: str, filename: str = "") -> list[str]:
    variants: list[str] = []
    for value in (title, filename):
        cleaned = clean_title_for_search(value)
        if cleaned:
            variants.append(cleaned)
        parsed_artist, parsed_title = parse_artist_title(value)
        if parsed_artist and parsed_title:
            variants.append(parsed_title)
    for value in list(variants):
        parts = value.split()
        if len(parts) >= 3 and (any(char.isdigit() for char in parts[0]) or parts[0].isupper()):
            variants.append(" ".join(parts[1:]))
    deduped: list[str] = []
    seen: set[str] = set()
    for variant in variants:
        key = normalize_text(variant)
        if key and key not in seen:
            seen.add(key)
            deduped.append(variant)
    return deduped[:5]


def _walk_json_artists(data: object) -> Iterable[str]:
    if isinstance(data, dict):
        for key, value in data.items():
            normalized_key = normalize_header(str(key))
            if normalized_key in ARTIST_JSON_FIELDS:
                if isinstance(value, list):
                    for item in value:
                        yield clean_text(item if not isinstance(item, dict) else item.get("name", ""))
                else:
                    yield clean_text(value)
            yield from _walk_json_artists(value)
    elif isinstance(data, list):
        for item in data:
            yield from _walk_json_artists(item)


class MissingArtistResolver:
    def __init__(
        self,
        *,
        media_roots: Sequence[Path | str] = (),
        search_endpoint: str | None = None,
        min_confidence: float | None = None,
        max_results: int | None = None,
        attempts: int | None = None,
        timeout: float | None = None,
        sleep_seconds: float | None = None,
        session: requests.Session | None = None,
        search_enabled: bool | None = None,
    ) -> None:
        self.media_roots = tuple(Path(root) for root in media_roots if str(root))
        self.search_endpoint = (
            search_endpoint
            or os.getenv("ARTIST_REPAIR_SEARCH_ENDPOINT")
            or os.getenv("ARTIST_SEARCH_ENDPOINT")
            or "https://www.google.com/search"
        )
        self.min_confidence = (
            float(min_confidence)
            if min_confidence is not None
            else float(os.getenv("ARTIST_REPAIR_MIN_CONFIDENCE", os.getenv("ARTIST_SEARCH_MIN_CONFIDENCE", "0.86")))
        )
        self.max_results = int(max_results if max_results is not None else os.getenv("ARTIST_REPAIR_MAX_RESULTS", "12"))
        self.attempts = int(attempts if attempts is not None else os.getenv("ARTIST_REPAIR_ATTEMPTS", "3"))
        self.timeout = float(timeout if timeout is not None else os.getenv("ARTIST_REPAIR_TIMEOUT_SECONDS", "12"))
        self.sleep_seconds = float(
            sleep_seconds if sleep_seconds is not None else os.getenv("ARTIST_REPAIR_SLEEP_SECONDS", "0.35")
        )
        self.search_enabled = (
            search_enabled
            if search_enabled is not None
            else os.getenv("ARTIST_REPAIR_SEARCH_ENABLED", "true").strip().lower() in {"1", "true", "yes", "on"}
        )
        self.session = session or requests.Session()
        self._cache: dict[str, ArtistResolution | None] = {}
        self._media_index: dict[str, Path] | None = None

    def resolve(
        self,
        *,
        title: str,
        filename: str = "",
        file_path: str = "",
        existing_artist: str = "",
        row: Mapping[str, object] | None = None,
        header_map: Mapping[str, str] | None = None,
        csv_context: str = "",
    ) -> ArtistResolution | None:
        if not is_missing_artist_value(existing_artist):
            return ArtistResolution(clean_text(existing_artist), 1.0, "existing")
        if is_missing_title_value(title):
            return None

        row = row or {}
        header_map = header_map or build_header_map(list(row.keys()))
        local = self._resolve_from_local_evidence(
            title=title,
            filename=filename,
            file_path=file_path,
            row=row,
            header_map=header_map,
            csv_context=csv_context,
        )
        if local and local.confidence >= self.min_confidence:
            return local
        google = self._resolve_from_google(title=title, filename=filename, row=row, header_map=header_map)
        if google and (not local or google.confidence >= local.confidence):
            return google
        return local if local and local.confidence >= self.min_confidence else None

    def _resolve_from_local_evidence(
        self,
        *,
        title: str,
        filename: str,
        file_path: str,
        row: Mapping[str, object],
        header_map: Mapping[str, str],
        csv_context: str,
    ) -> ArtistResolution | None:
        parsed_artist, parsed_title = parse_artist_title(filename or file_path)
        if parsed_artist and self._title_matches(title, parsed_title):
            return ArtistResolution(parsed_artist, 0.96, "local_filename", filename or file_path)

        for column in PROVIDER_ARTIST_COLUMNS:
            candidate = _artist_candidate(row_value(row, header_map, column), title=title, parsed_title=parsed_title)
            if candidate:
                return ArtistResolution(candidate, 0.92, f"local_column:{column}")

        for column in PROVIDER_JSON_COLUMNS:
            raw = row_value(row, header_map, column)
            if not raw:
                continue
            for candidate in self._artists_from_json(raw, title=title, parsed_title=parsed_title):
                return ArtistResolution(candidate, 0.90, f"local_json:{column}")

        media_path = self._resolve_media_path(file_path or filename, filename, csv_context)
        if media_path:
            embedded = self._embedded_artist(media_path, title=title, parsed_title=parsed_title)
            if embedded:
                return ArtistResolution(embedded, 0.95, "local_embedded_tag", str(media_path))
        return None

    def _artists_from_json(self, raw: str, *, title: str, parsed_title: str) -> Iterable[str]:
        try:
            data = json.loads(raw)
        except (TypeError, ValueError):
            return []
        candidates = []
        for value in _walk_json_artists(data):
            candidate = _artist_candidate(value, title=title, parsed_title=parsed_title)
            if candidate:
                candidates.append(candidate)
        return candidates

    def _embedded_artist(self, path: Path, *, title: str, parsed_title: str) -> str:
        try:
            media = MutagenFile(path, easy=True)
        except Exception:
            return ""
        if not media or not getattr(media, "tags", None):
            return ""
        for key in ("artist", "albumartist", "performer", "composer"):
            values = media.tags.get(key) if media.tags else None
            if not values:
                continue
            if isinstance(values, str):
                values = [values]
            for value in values:
                candidate = _artist_candidate(value, title=title, parsed_title=parsed_title)
                if candidate:
                    return candidate
        return ""

    def _resolve_media_path(self, file_path: str, filename: str, csv_context: str) -> Path | None:
        values = [clean_text(file_path), clean_text(filename)]
        for value in values:
            if not value:
                continue
            candidate = Path(value)
            if candidate.is_file():
                return candidate
            relative = value.replace("\\", "/")
            relative = re.sub(r"^/app/input_media/?", "", relative)
            for root in self.media_roots:
                for joined in (root / relative, root / csv_context / relative, root / Path(relative).name):
                    if joined.is_file():
                        return joined
        if filename:
            return self._media_index_lookup(filename)
        return None

    def _media_index_lookup(self, filename: str) -> Path | None:
        if self._media_index is None:
            index: dict[str, Path] = {}
            for root in self.media_roots:
                if not root.exists():
                    continue
                for path in root.rglob("*"):
                    if path.is_file() and path.suffix.lower() in MEDIA_EXTENSIONS:
                        index.setdefault(path.name.casefold(), path)
            self._media_index = index
        return self._media_index.get(Path(filename).name.casefold())

    def _title_matches(self, title: str, candidate_title: str) -> bool:
        if not title or not candidate_title:
            return False
        expected = normalize_text(clean_title_for_search(title))
        candidate = normalize_text(clean_title_for_search(candidate_title))
        return bool(expected and candidate and (expected == candidate or similarity(expected, candidate) >= 0.82))

    def _resolve_from_google(
        self,
        *,
        title: str,
        filename: str,
        row: Mapping[str, object],
        header_map: Mapping[str, str],
    ) -> ArtistResolution | None:
        if not self.search_enabled:
            return None
        best: ArtistResolution | None = None
        for query in self._queries(title, filename, row, header_map):
            if query in self._cache:
                cached = self._cache[query]
                if cached and (best is None or cached.confidence > best.confidence):
                    best = cached
                continue
            resolution = self._search_once(query, title=title, filename=filename)
            self._cache[query] = resolution
            if resolution and (best is None or resolution.confidence > best.confidence):
                best = resolution
            if best and best.confidence >= 0.94:
                break
        return best if best and best.confidence >= self.min_confidence else None

    def _queries(
        self,
        title: str,
        filename: str,
        row: Mapping[str, object],
        header_map: Mapping[str, str],
    ) -> list[str]:
        year = row_value(row, header_map, "year", "release_year")
        album = row_value(row, header_map, "album")
        genre = row_value(row, header_map, "genre")
        queries: list[str] = []
        for variant in _title_variants(title, filename):
            queries.extend(
                [
                    f'the artist for the song "{variant}"',
                    f'"{variant}" artist song',
                    f'"{variant}" official song artist',
                    f'"{variant}" site:music.apple.com',
                    f'"{variant}" site:open.spotify.com',
                    f'"{variant}" site:musicbrainz.org',
                ]
            )
            if year:
                queries.append(f'"{variant}" "{year}" artist song')
            if album and normalize_text(album) not in {"single", "unknown"}:
                queries.append(f'"{variant}" "{album}" artist')
            if genre:
                queries.append(f'"{variant}" "{genre}" artist song')
        if filename:
            queries.append(f'"{_strip_extension_noise(filename)}" artist song')
        deduped: list[str] = []
        seen: set[str] = set()
        for query in queries:
            key = normalize_text(query)
            if key and key not in seen:
                seen.add(key)
                deduped.append(query)
        return deduped[: self.max_results]

    def _search_once(self, query: str, *, title: str, filename: str) -> ArtistResolution | None:
        params = {"q": query, "hl": "en", "num": "10"}
        headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"
            )
        }
        last_error = ""
        for attempt in range(max(1, self.attempts)):
            try:
                response = self.session.get(
                    self.search_endpoint,
                    params=params,
                    headers=headers,
                    timeout=self.timeout,
                )
                if response.status_code in {429, 500, 502, 503, 504}:
                    last_error = f"status={response.status_code}"
                    time.sleep(self.sleep_seconds * (attempt + 1))
                    continue
                response.raise_for_status()
                return self._parse_google_html(response.text, title=title, filename=filename, query=query)
            except requests.RequestException as exc:
                last_error = str(exc)
                time.sleep(self.sleep_seconds * (attempt + 1))
        return ArtistResolution("", 0.0, "google_failed", last_error) if last_error else None

    def _parse_google_html(self, html: str, *, title: str, filename: str, query: str) -> ArtistResolution | None:
        soup = BeautifulSoup(html, "html.parser")
        blocks = []
        for selector in SEARCH_RESULT_SELECTORS:
            blocks.extend(soup.select(selector))
        if not blocks:
            blocks = [soup]

        votes: dict[str, tuple[float, str, int]] = {}
        title_variants = _title_variants(title, filename)
        for index, block in enumerate(blocks[:12]):
            text = clean_text(block.get_text(" ", strip=True))
            if not text:
                continue
            score_base = 0.70 + max(0, 8 - index) * 0.015
            lower_text = text.lower()
            if any(hint in lower_text for hint in TRUSTED_SOURCE_HINTS):
                score_base += 0.08
            for variant in title_variants:
                for candidate, pattern_score in self._candidates_from_text(text, variant):
                    clean_candidate = _artist_candidate(candidate, title=title, parsed_title=variant)
                    if not clean_candidate:
                        continue
                    score = min(0.98, score_base + pattern_score)
                    key = normalize_text(clean_candidate)
                    current = votes.get(key)
                    count = (current[2] + 1) if current else 1
                    if current:
                        score = max(score, current[0]) + min(0.06, count * 0.015)
                    votes[key] = (min(score, 0.99), clean_candidate, count)

        if not votes:
            return None
        confidence, artist, count = max(votes.values(), key=lambda item: (item[0], item[2], len(item[1])))
        if confidence < self.min_confidence:
            return None
        return ArtistResolution(artist, confidence, "google_search", query)

    def _candidates_from_text(self, text: str, title: str) -> Iterable[tuple[str, float]]:
        title_re = re.escape(title)
        patterns = [
            (rf"artist\s+for\s+the\s+song\s+['\"]?{title_re}['\"]?\s+is\s+([^.\n|]+)", 0.20),
            (rf"{title_re}\s+(?:is\s+)?(?:a\s+song\s+)?by\s+([^.\n|]+)", 0.18),
            (rf"{title_re}\s*[-|]\s*(?:song\s+by\s+)?([^.\n|]+)", 0.16),
            (rf"([^.\n|:-]{{2,80}})\s*[-|]\s*{title_re}", 0.16),
            (rf"{title_re}.*?\bartist\s*[:\-]\s*([^.\n|]+)", 0.14),
            (rf"\bartist\s*[:\-]\s*([^.\n|]+).*?{title_re}", 0.14),
            (rf"\bkaraoke\s+{title_re}\s*[-|]\s*([^.\n|]+)", 0.12),
        ]
        for pattern, score in patterns:
            for match in re.finditer(pattern, text, flags=re.IGNORECASE):
                yield match.group(1), score


def repair_rows(
    rows: Sequence[dict[str, str]],
    headers: list[str],
    *,
    resolver: MissingArtistResolver,
    media_roots: Sequence[Path | str] = (),
    csv_context: str = "",
    remove_unresolved: bool = True,
) -> tuple[list[dict[str, str]], list[str], ArtistRepairStats]:
    header_map = build_header_map(headers)
    if "artist" not in header_map:
        headers.append("artist")
        header_map = build_header_map(headers)
    repaired_rows: list[dict[str, str]] = []
    stats = ArtistRepairStats()
    for row in rows:
        stats.checked += 1
        title = row_value(row, header_map, "title")
        artist_key = header_map.get("artist", "artist")
        artist = row_value(row, header_map, "artist")
        if is_missing_title_value(title) and is_missing_artist_value(artist):
            stats.removed_missing_identity += 1
            continue
        if not is_missing_artist_value(artist):
            stats.unchanged += 1
            repaired_rows.append(row)
            continue

        filename = filename_from_row(row, header_map)
        file_path = row_value(row, header_map, "file_path", "path", "filepath")
        resolution = resolver.resolve(
            title=title,
            filename=filename,
            file_path=file_path,
            existing_artist=artist,
            row=row,
            header_map=header_map,
            csv_context=csv_context,
        )
        if resolution and resolution.artist:
            row = dict(row)
            row[artist_key] = resolution.artist
            stats.filled += 1
            if resolution.source.startswith("google"):
                stats.google_hits += 1
            else:
                stats.local_hits += 1
            repaired_rows.append(row)
            continue

        if remove_unresolved:
            stats.removed_unresolved_artist += 1
        else:
            repaired_rows.append(row)
    return repaired_rows, headers, stats


def tabular_cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def read_tabular_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=False)
        sheet = workbook[workbook.sheetnames[0]]
        raw_rows = list(sheet.iter_rows(values_only=True))
        if not raw_rows:
            return [], []
        headers = [tabular_cell_text(value) for value in raw_rows[0]]
        rows = []
        for raw in raw_rows[1:]:
            row = {
                headers[index]: tabular_cell_text(raw[index] if index < len(raw) else "")
                for index in range(len(headers))
                if headers[index]
            }
            if any(row.values()):
                rows.append(row)
        return headers, rows

    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        return list(reader.fieldnames or []), list(reader)


def write_tabular_rows(path: Path, headers: Sequence[str], rows: Sequence[Mapping[str, str]]) -> None:
    if path.suffix.lower() == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "WithTag"
        sheet.append(list(headers))
        for row in rows:
            sheet.append([tabular_cell_text(row.get(header, "")) for header in headers])
        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        for index, header in enumerate(headers, start=1):
            values = [str(header), *(str(row.get(header, "")) for row in rows[:200])]
            width = min(max(max((len(value) for value in values), default=10) + 2, 10), 48)
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        handle = tempfile.NamedTemporaryFile(delete=False, dir=str(path.parent), suffix=".xlsx")
        temp_path = Path(handle.name)
        handle.close()
        try:
            workbook.save(temp_path)
            temp_path.replace(path)
        finally:
            if temp_path.exists():
                temp_path.unlink()
        return

    handle = tempfile.NamedTemporaryFile(
        "w",
        newline="",
        encoding="utf-8-sig",
        delete=False,
        dir=str(path.parent),
        prefix=f".{path.name}.",
        suffix=".part",
    )
    temp_path = Path(handle.name)
    try:
        with handle:
            writer = csv.DictWriter(handle, fieldnames=list(headers), extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)
        temp_path.replace(path)
    finally:
        if temp_path.exists():
            temp_path.unlink()
